# -*- coding: utf-8 -*-
"""Carga la tabla `dter` desde los crudos CND (DTER Oficial, hoja DATOS INFORMATIVOS).

Es la tabla que leia el sistema viejo para la linea "DTER de <mes>" de la
factura (funcion fn_DTER_Mes):

    fn_DTER_Mes(@cCliente, @nAnnio, @nMes) =
        SUM(IIF(Atributo='Debito', Valor, -Valor))  FROM DTER
        WHERE Cliente=@cCliente AND Annio=@nAnnio AND Mes=@nMes

El DTER de un periodo se arma con el archivo DTER Oficial de ese periodo.
La factura de un mes incluye el DTER del mes anterior (igual que el sistema viejo).

Uso:
    python importar_dter.py                    # todos los periodos presentes
    python importar_dter.py --periodo 202606
    python importar_dter.py --listar
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
sys.path.insert(0, str(Path(__file__).resolve().parent))

from sqlalchemy import text  # noqa: E402

from app.core.database import engine  # noqa: E402

BASE = Path(__file__).resolve().parent.parent
RAW = BASE / "CND_Scraper" / "raw"

DDL = """
CREATE TABLE IF NOT EXISTS dter (
    id BIGINT NOT NULL AUTO_INCREMENT,
    annio INT DEFAULT NULL,
    mes INT DEFAULT NULL,
    generador VARCHAR(20) DEFAULT NULL,
    cliente VARCHAR(20) DEFAULT NULL,
    transaccion VARCHAR(100) DEFAULT NULL,
    atributo VARCHAR(10) DEFAULT NULL,
    valor DECIMAL(9,2) DEFAULT NULL,
    periodo CHAR(6) GENERATED ALWAYS AS (CONCAT(annio, LPAD(mes,2,'0'))) STORED,
    archivo_origen VARCHAR(255) DEFAULT NULL,
    fecha_carga DATETIME DEFAULT CURRENT_TIMESTAMP,
    PRIMARY KEY (id),
    UNIQUE KEY uk_dter (cliente, generador, periodo, transaccion, atributo),
    KEY idx_dter_periodo (periodo),
    KEY idx_dter_cliente (cliente)
) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;
"""


def archivos_dter() -> list[tuple[str, str, Path]]:
    out = []
    for carpeta in RAW.rglob("CND_DTER_OFICIAL"):
        generador = carpeta.parent.name
        for sub in sorted(p for p in carpeta.iterdir() if p.is_dir()):
            periodo = sub.name
            if not (periodo.isdigit() and len(periodo) == 6):
                continue
            for arch in sorted(sub.glob("*.xlsx")):
                out.append((generador, periodo, arch))
    return out


def atributo_de(valor: str) -> str:
    v = (valor or "").strip().lower()
    if v.startswith("d"):
        return "Debito"
    if v.startswith("c"):
        return "Credito"
    return ""


def leer_datos_informativos(arch: Path) -> list[dict]:
    """Devuelve los renglones de la hoja DATOS INFORMATIVOS."""
    import openpyxl

    wb = openpyxl.load_workbook(arch, read_only=True, data_only=True)
    if "DATOS INFORMATIVOS" not in wb.sheetnames:
        wb.close()
        return []
    ws = wb["DATOS INFORMATIVOS"]
    filas = list(ws.iter_rows(values_only=True))
    wb.close()

    # localiza el header
    ini = None
    for i, fila in enumerate(filas):
        c0 = str(fila[0] or "").strip().lower()
        if c0.startswith("c") and "agente" in c0:
            ini = i + 1
            break
    if ini is None:
        return []

    out = []
    for fila in filas[ini:]:
        cliente = str(fila[0] or "").strip()
        if not cliente:
            continue
        trans = str(fila[1] or "").strip()
        atributo = atributo_de(str(fila[2] or ""))
        try:
            valor = float(fila[3] or 0)
        except (TypeError, ValueError):
            valor = 0.0
        if not atributo or abs(valor) < 0.005:
            continue
        out.append({"cliente": cliente, "transaccion": trans,
                    "atributo": atributo, "valor": round(valor, 2)})
    return out


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--periodo")
    ap.add_argument("--listar", action="store_true")
    args = ap.parse_args()

    todos = archivos_dter()
    if args.periodo:
        todos = [t for t in todos if t[1] == args.periodo]

    with engine.begin() as cn:
        if not args.listar:
            cn.execute(text(DDL))

        gran_total = 0
        for generador, periodo, arch in sorted(todos, key=lambda t: (t[1], t[2].name)):
            filas = leer_datos_informativos(arch)
            if not filas:
                print(f"    {generador}/{periodo}: {arch.name} (sin DATOS INFORMATIVOS)")
                continue
            if not args.listar:
                cn.execute(text("DELETE FROM dter WHERE generador=:g AND periodo=:p"),
                           {"g": generador, "p": periodo})
                cn.execute(text(
                    "INSERT INTO dter (annio, mes, generador, cliente, transaccion, atributo, valor, archivo_origen) "
                    "VALUES (:annio,:mes,:generador,:cliente,:transaccion,:atributo,:valor,:archivo_origen) "
                    "ON DUPLICATE KEY UPDATE valor=VALUES(valor), archivo_origen=VALUES(archivo_origen)"
                ), [{**f, "annio": int(periodo[:4]), "mes": int(periodo[4:6]),
                     "generador": generador, "archivo_origen": arch.name} for f in filas])
            gran_total += len(filas)
            suma = sum(f["valor"] if f["atributo"] == "Debito" else -f["valor"] for f in filas)
            print(f"    {generador}/{periodo}: {arch.name} -> {len(filas)} renglones (neto {suma:.2f})")

        print(f"[ok] {gran_total} renglones" + (" (listar)" if args.listar else ""))


if __name__ == "__main__":
    main()
