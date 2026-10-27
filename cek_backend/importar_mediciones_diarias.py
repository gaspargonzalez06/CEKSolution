# -*- coding: utf-8 -*-
from __future__ import annotations
import re
import sys
from pathlib import Path
from collections import OrderedDict
from datetime import datetime

sys.path.insert(0, str(Path(__file__).resolve().parent))
from sqlalchemy import text
from app.core.database import engine

BASE = Path(__file__).resolve().parent.parent
RAW = BASE / "CND_Scraper" / "raw"

DDL = """
CREATE TABLE IF NOT EXISTS mediciones_diarias (
  id BIGINT NOT NULL AUTO_INCREMENT,
  annio INT NOT NULL,
  mes INT NOT NULL,
  periodo CHAR(6) NOT NULL,
  fecha DATE NOT NULL,
  generador VARCHAR(20) NOT NULL,
  cliente VARCHAR(20) NOT NULL,
  consumo DECIMAL(14,6) NOT NULL DEFAULT 0,
  consumo_kwh DECIMAL(18,3) NOT NULL DEFAULT 0,
  cms_prom DECIMAL(12,4) NULL,
  archivo_origen VARCHAR(255) NOT NULL,
  fecha_carga DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP,
  PRIMARY KEY (id),
  UNIQUE KEY uk_md (cliente, generador, fecha),
  KEY idx_md_periodo (periodo),
  KEY idx_md_cliente (cliente),
  KEY idx_md_fecha (fecha)
) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;
"""

RE_CLIENTE = re.compile(r"_(gcp\d{4})_", re.I)


def archivos_liq() -> list[tuple[str, str, str, Path]]:
    out = []
    for carpeta in RAW.rglob("CND_DTE_LIQUIDACION"):
        generador = carpeta.parent.name
        for sub in sorted(p for p in carpeta.iterdir() if p.is_dir()):
            periodo = sub.name
            if len(periodo) != 6 or not periodo.isdigit():
                continue
            for arch in sorted(sub.glob("*.xlsx")):
                if "_aj" in arch.name.lower():
                    continue
                m = RE_CLIENTE.search(arch.name)
                if not m:
                    continue
                out.append((generador, periodo, m.group(1).upper(), arch))
    return out


def leer_diario(arch: Path, cliente: str) -> list[dict]:
    import openpyxl

    wb = openpyxl.load_workbook(arch, read_only=True, data_only=True)
    if "MEDICIONES" not in wb.sheetnames:
        wb.close()
        return []
    ws = wb["MEDICIONES"]
    filas = list(ws.iter_rows(values_only=True))
    wb.close()

    # header
    hdr = None
    for i, f in enumerate(filas[:15]):
        celdas = [str(c).strip().lower() if c is not None else "" for c in f]
        if "fecha" in celdas and ("hora" in celdas or "consumo" in celdas):
            hdr = (i, [str(c).strip().lower() if c is not None else "" for c in f])
            break
    if not hdr:
        return []
    idx, h = hdr
    i_fecha = h.index("fecha")
    i_hora = h.index("hora") if "hora" in h else 1
    i_cms = h.index("cms") if "cms" in h else None
    i_sub = h.index("subsistema") if "subsistema" in h else None

    # buscar columna de consumo del cliente: brillosol_a_e? o columna 'consumo' global? 
    # pero por cliente: en archivos per-cliente, puede variar; buscar columna con 'consumo' o con sufijo
    col_cons = None
    # intentar 'consumo' exacto
    if "consumo" in h:
        col_cons = h.index("consumo")
    else:
        # buscar columna que contenga _a_e del cliente? raro en formato; tomar la columna mas 'consumo-like'
        for j, c in enumerate(h):
            if "_a_e" in c or "_cons" in c:
                col_cons = j
                break
    if col_cons is None:
        # fallback: columna que no sea metadatos
        for j, c in enumerate(h):
            if j not in (i_fecha, i_hora, i_cms, i_sub) and c:
                col_cons = j
                break
    if col_cons is None:
        return []

    # agrupar por fecha: consumo sum (MWh), cms promedio (de horas)
    by_fecha = OrderedDict()
    for f in filas[idx + 1 :]:
        if f[i_fecha] is None:
            continue
        fecha = f[i_fecha]
        if hasattr(fecha, "date"):
            fecha = fecha.date()
        fecha_s = fecha.isoformat()
        cms_v = f[i_cms] if i_cms is not None and i_cms < len(f) else None
        cons_v = f[col_cons] if col_cons < len(f) else None
        e = by_fecha.setdefault(fecha_s, {"cms": [], "cons": 0.0})
        if isinstance(cms_v, (int, float)):
            e["cms"].append(float(cms_v))
        if isinstance(cons_v, (int, float)):
            e["cons"] += float(cons_v)

    out = []
    for fs, e in by_fecha.items():
        fdt = datetime.strptime(fs, "%Y-%m-%d").date()
        cons_mwh = float(e["cons"] or 0)
        cms = None
        if e["cms"]:
            cms = round(sum(e["cms"]) / len(e["cms"]), 4)
        out.append(
            {
                "annio": fdt.year,
                "mes": fdt.month,
                "periodo": f"{fdt.year:04d}{fdt.month:02d}",
                "fecha": fs,
                "consumo": round(cons_mwh, 6),
                "consumo_kwh": round(cons_mwh * 1000, 3),
                "cms_prom": cms,
            }
        )
    return out


def main() -> None:
    import argparse

    ap = argparse.ArgumentParser()
    ap.add_argument("--periodo")
    ap.add_argument("--listar", action="store_true")
    args = ap.parse_args()

    todos = archivos_liq()
    if args.periodo:
        todos = [t for t in todos if t[1] == args.periodo]

    with engine.begin() as cn:
        if not args.listar:
            cn.execute(text(DDL))
        filas = []
        for generador, periodo, cliente, arch in sorted(todos, key=lambda t: (t[1], t[2])):
            diario = leer_diario(arch, cliente)
            if not diario:
                print(f"  {generador}/{periodo} {cliente}: sin datos ({arch.name})")
                continue
            for d in diario:
                filas.append(
                    {
                        **d,
                        "generador": generador,
                        "cliente": cliente,
                        "archivo_origen": arch.name,
                    }
                )
            print(f"  {generador}/{periodo} {cliente}: {len(diario)} días ({arch.name})")
        if filas and not args.listar:
            per_set = {f"{f['annio']}{f['mes']:02d}" for f in filas}
            for p in per_set:
                cn.execute(text("DELETE FROM mediciones_diarias WHERE periodo=:p"), {"p": p})
            cn.execute(
                text(
                    "INSERT INTO mediciones_diarias (annio, mes, periodo, fecha, generador, cliente, consumo, consumo_kwh, cms_prom, archivo_origen) "
                    "VALUES (:annio,:mes,:periodo,:fecha,:generador,:cliente,:consumo,:consumo_kwh,:cms_prom,:archivo_origen)"
                ),
                filas,
            )
        print(f"[ok] {len(filas)} registros diarios" + (" (listar)" if args.listar else ""))


if __name__ == "__main__":
    main()
